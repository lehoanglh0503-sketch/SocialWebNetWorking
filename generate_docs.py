import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document(docx_path):
    doc = Document()
    
    # Set page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("BÁO CÁO BÀI TẬP LỚN / ĐỒ ÁN MÔN HỌC\n")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    run_main = title_p.add_run("PHÂN TÍCH YÊU CẦU PHẦN MỀM HỆ THỐNG SOCIAL WEBSITE (CONNECTIFY)")
    run_main.font.name = 'Times New Roman'
    run_main.font.size = Pt(16)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_after = Pt(14)
    run_tech = meta_p.add_run("Công nghệ: ASP.NET Core 6.0 MVC | Entity Framework Core | SQL Server | SignalR | Bootstrap 5")
    run_tech.font.name = 'Times New Roman'
    run_tech.font.size = Pt(11)
    run_tech.font.italic = True
    run_tech.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Member list table
    p_mem = doc.add_paragraph()
    p_mem.paragraph_format.space_before = Pt(6)
    p_mem.paragraph_format.space_after = Pt(4)
    run_mem_title = p_mem.add_run("DANH SÁCH THÀNH VIÊN VÀ ĐÁNH GIÁ MỨC ĐỘ ĐÓNG GÓP")
    run_mem_title.bold = True
    run_mem_title.font.size = Pt(12)
    run_mem_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    mem_table = doc.add_table(rows=4, cols=6)
    mem_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(mem_table)
    
    col_widths = [Inches(0.6), Inches(1.8), Inches(1.1), Inches(1.1), Inches(1.9), Inches(0.9)]
    headers = ["STT", "Họ và tên", "Mã sinh viên", "Vai trò", "Nhiệm vụ chính", "Đóng góp"]
    
    # Format Header Row
    hdr_cells = mem_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10.5)

    members_data = [
        ("1", "[Họ và tên bạn]", "[MSSV 1]", "Nhóm trưởng", "Phân tích yêu cầu, thiết kế CSDL, dựng Backend Auth, Post, SignalR", "100%"),
        ("2", "[Thành viên 2]", "[MSSV 2]", "Thành viên", "Xây dựng tính năng Chat thời gian thực, quản lý Bạn bè (Friend, Chat)", "100%"),
        ("3", "[Thành viên 3]", "[MSSV 3]", "Thành viên", "Dashboard Quản trị (Admin), kiểm duyệt vi phạm, quản lý từ cấm, UI", "100%")
    ]

    for row_idx, data in enumerate(members_data, start=1):
        row_cells = mem_table.rows[row_idx].cells
        bg_color = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 2, 3, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper function for adding headings
    def add_sec_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return h

    def add_sub_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2B, 0x4C, 0x7E)
        return h

    def add_bullet(p, bold_prefix, content):
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        run_b = p.add_run(bold_prefix + " ")
        run_b.bold = True
        run_b.font.size = Pt(11.5)
        run_c = p.add_run(content)
        run_c.font.size = Pt(11.5)

    # -------------------------------------------------------------
    # a. MỤC TIÊU PHẦN MỀM
    # -------------------------------------------------------------
    add_sec_heading("a. Mục tiêu phần mềm")
    
    p = doc.add_paragraph()
    add_bullet(p, "• Nền tảng mạng xã hội kết nối trực tuyến:", 
               "Xây dựng website mạng xã hội hiện đại giúp người dùng đã xác thực chia sẻ bài viết dạng văn bản và hình ảnh đa phương tiện; tương tác thời gian thực thông qua hệ thống cảm xúc phong phú (6 loại biểu tượng cảm xúc) và bình luận phân cấp lồng nhau (threaded comments).")

    p = doc.add_paragraph()
    add_bullet(p, "• Tương tác và giao tiếp thời gian thực (Real-time):", 
               "Tích hợp công nghệ SignalR (SocialHub) để tự động phát sóng bài viết mới, cập nhật số lượng cảm xúc, phản hồi bình luận và gửi tin nhắn tức thời (bao gồm tin nhắn chữ, ảnh tải lên và thẻ bài viết chia sẻ trực tiếp) giữa bạn bè mà không yêu cầu người dùng phải tải lại trang (reload).")

    p = doc.add_paragraph()
    add_bullet(p, "• Cung cấp bộ công cụ quản trị (Admin Dashboard) toàn diện:", 
               "Trang bị công cụ quản trị cho người quản lý hệ thống xem thống kê tổng quan (người dùng, bài viết, báo cáo); quản lý danh sách tài khoản (đặt lại mật khẩu, khóa/mở khóa tài khoản, xóa tài khoản cascade); kiểm duyệt nội dung bài viết và quản trị từ điển từ cấm.")

    p = doc.add_paragraph()
    add_bullet(p, "• Kiểm soát nội dung và bảo vệ cộng đồng tự động:", 
               "Tự động phát hiện và ngăn chặn nội dung chứa từ ngữ bạo lực, phản cảm thông qua danh sách từ cấm (BannedWords) bằng biểu thức chính quy (Regex); thiết lập quy chế xử lý vi phạm bậc thang tự động (cảnh báo vi phạm ở lần 2 và tự động khóa tài khoản ở lần thứ 3 khi Admin phê duyệt báo cáo vi phạm).")

    p = doc.add_paragraph()
    add_bullet(p, "• Kiến trúc lưu trữ và khởi tạo hệ thống chuẩn mực:", 
               "Lưu trữ dữ liệu an toàn và nhất quán trên cơ sở dữ liệu quan hệ Microsoft SQL Server thông qua Entity Framework Core (SocialDbContext), hỗ trợ cơ chế tự động khởi tạo dữ liệu mẫu (SeedData) phục vụ trải nghiệm và kiểm thử hệ thống ngay khi triển khai.")

    # -------------------------------------------------------------
    # b. CÁC ROLE (VAI TRÒ NGƯỜI DÙNG)
    # -------------------------------------------------------------
    add_sec_heading("b. Các role trong hệ thống")

    role_table = doc.add_table(rows=4, cols=2)
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(role_table)
    
    role_hdr = role_table.rows[0].cells
    role_hdr[0].text = "Role (Vai trò)"
    role_hdr[1].text = "Mô tả & Phạm vi quyền hạn chi tiết"
    for c in role_hdr:
        set_cell_background(c, "1B365D")
        set_cell_margins(c, top=120, bottom=120, left=120, right=120)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(11)

    roles_data = [
        ("Khách chưa đăng nhập\n(Guest)",
         "• Truy cập trang giới thiệu (/Home/About) và trang chính sách bảo mật (/Home/Privacy).\n"
         "• Đăng ký tài khoản người dùng mới (/Auth/Register).\n"
         "• Đăng nhập vào hệ thống (/Auth/Login).\n"
         "• Bị chặn truy cập các chức năng tương tác nội bộ: khi cố truy cập Bảng tin, Trang cá nhân, Bạn bè sẽ tự động được chuyển hướng về trang Đăng nhập."),
        
        ("Người dùng đã đăng nhập\n(Authenticated User)",
         "• Xem bảng tin các bài viết mới nhất; đăng bài viết (văn bản + ảnh); xóa bài viết của chính mình; chia sẻ bài viết qua tin nhắn.\n"
         "• Bày tỏ 6 loại cảm xúc trên bài viết và bình luận; xem danh sách những người đã thả cảm xúc.\n"
         "• Viết bình luận gốc và trả lời bình luận lồng nhau; xóa bình luận của mình hoặc trên bài viết của mình.\n"
         "• Xem đề xuất kết bạn; gửi, chấp nhận, từ chối lời mời hoặc hủy kết bạn.\n"
         "• Nhắn tin 1-1 với bạn bè (gửi văn bản, gửi file ảnh, gửi bài viết chia sẻ) theo thời gian thực (SignalR).\n"
         "• Tìm kiếm người dùng và bài viết theo từ khóa.\n"
         "• Báo cáo bài viết vi phạm kèm lý do chi tiết.\n"
         "• Quản lý hồ sơ cá nhân: xem thống kê cá nhân, đổi họ tên, đổi ảnh đại diện (avatar), đổi mật khẩu."),
        
        ("Quản trị viên\n(Admin / Moderator)",
         "• Tự động chuyển hướng vào Dashboard quản trị (/Admin/Index) khi đăng nhập.\n"
         "• Xem thống kê hệ thống: Tổng số người dùng, tổng số bài viết, tổng số báo cáo vi phạm.\n"
         "• Quản lý người dùng: Xem danh sách thành viên, đặt lại mật khẩu, khóa/mở khóa tài khoản, xóa tài khoản vĩnh viễn (xóa cascade toàn bộ dữ liệu liên quan).\n"
         "• Kiểm duyệt nội dung: Xem toàn bộ bài viết, thực hiện xóa bài viết vi phạm.\n"
         "• Xử lý báo cáo vi phạm: Bỏ qua báo cáo hoặc Duyệt báo cáo (xóa bài viết, cộng 1 lần vi phạm cho tác giả, tự động khóa tài khoản khi số lần vi phạm >= 3).\n"
         "• Quản lý từ cấm: Xem danh sách, thêm từ cấm mới, xóa từ cấm khỏi hệ thống.")
    ]

    for idx, (role_name, desc) in enumerate(roles_data, start=1):
        cells = role_table.rows[idx].cells
        cells[0].text = role_name
        cells[1].text = desc
        bg = "F7F9FC" if idx % 2 == 1 else "FFFFFF"
        for c_idx, c in enumerate(cells):
            set_cell_background(c, bg)
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            p = c.paragraphs[0]
            if c_idx == 0:
                p.runs[0].font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # c. CÁC CHỨC NĂNG CƠ BẢN VÀ CHI TIẾT
    # -------------------------------------------------------------
    add_sec_heading("c. Các chức năng cơ bản và chi tiết")

    func_table = doc.add_table(rows=12, cols=3)
    func_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(func_table)

    f_hdr = func_table.rows[0].cells
    f_hdr[0].text = "Nhóm chức năng"
    f_hdr[1].text = "Chức năng chi tiết"
    f_hdr[2].text = "Mô tả kỹ thuật & Hành vi hệ thống"
    for c in f_hdr:
        set_cell_background(c, "1B365D")
        set_cell_margins(c, top=120, bottom=120, left=100, right=100)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10.5)

    func_data = [
        ("1. Xác thực & Session", "Đăng ký tài khoản", "Nhập Username (3-50 ký tự), Họ tên, Mật khẩu (>=6 ký tự) và Xác nhận mật khẩu. Kiểm tra trùng tài khoản, băm SHA-256, tự sinh Avatar qua Dicebear API, tự động đăng nhập."),
        ("1. Xác thực & Session", "Đăng nhập & Đăng xuất", "Kiểm tra mật khẩu hash SHA-256 và cờ IsLocked. Nếu có 2 lần vi phạm đã duyệt thì hiển thị cảnh báo đỏ. Lưu Session (timeout 2h, cookie HttpOnly, IsEssential). Đăng xuất xóa toàn bộ Session."),
        ("2. Quản lý hồ sơ", "Xem & Cập nhật hồ sơ", "Xem thống kê cá nhân (số bài viết, số bình luận, số bạn bè). Cập nhật họ tên, tải ảnh đại diện mới (.jpg, .jpeg, .png, .gif, .svg) lưu vào /wwwroot/uploads/avatars/ với tên GUID."),
        ("2. Quản lý hồ sơ", "Đổi mật khẩu", "Nhập mật khẩu hiện tại (so sánh hash SHA-256), mật khẩu mới (tối thiểu 6 ký tự) và xác nhận mật khẩu mới. Lưu mật khẩu băm mới vào DB."),
        ("3. Quản lý bài viết", "Đăng bài viết mới", "Hỗ trợ bài viết văn bản (tối đa 1000 ký tự), hình ảnh (lưu tại /wwwroot/uploads/posts/) hoặc cả hai. Quét từ cấm bằng Regex. Tự động broadcast SignalR (ReceiveNewPost)."),
        ("3. Quản lý bài viết", "Xóa & Chia sẻ bài", "Tác giả bài viết hoặc Admin được xóa bài (broadcast SignalR ReceiveDeletedPost). Hỗ trợ nút Chia sẻ bài viết trực tiếp vào khung chat với bạn bè."),
        ("4. Bình luận & Phản hồi", "Bình luận lồng nhau", "Viết bình luận bài viết (tối đa 500 ký tự) và trả lời bình luận (Replies qua ParentCommentId). Kiểm tra từ cấm. Broadcast SignalR ReceiveNewComment."),
        ("5. Bày tỏ cảm xúc", "6 loại Reaction", "Thả cảm xúc trên bài viết và bình luận: Like, Love, Haha, Wow, Sad, Angry. Chọn lại cùng loại để gỡ, chọn loại khác để cập nhật. Đếm lượt và lấy Top 3 emoji phổ biến nhất. Realtime qua SignalR."),
        ("6. Bạn bè & Chat", "Quản lý bạn bè & Chat realtime", "Gợi ý kết bạn (ViewComponent loại trừ admin và bạn bè); gửi/chấp nhận/hủy kết bạn. Nhắn tin 1-1 với bạn bè (văn bản, gửi file ảnh, chia sẻ bài viết) qua SignalR (ReceiveDirectMessage)."),
        ("7. Tìm kiếm & Báo cáo", "Tìm kiếm & Báo cáo bài viết", "Tìm kiếm Người dùng (kèm trạng thái kết bạn) và Bài viết (tối đa 20 kết quả). Gửi báo cáo bài viết vi phạm kèm lý do (tối đa 500 ký tự, chống gửi trùng)."),
        ("8. Dashboard Admin", "Quản trị người dùng & Kiểm duyệt", "Thống kê tổng quan; quản lý thành viên (đổi pass, khóa/mở khóa, xóa cascade); duyệt báo cáo (xóa bài, tăng lần vi phạm, tự khóa khi đạt 3) hoặc bỏ qua; quản lý từ điển từ cấm (thêm/xóa).")
    ]

    for idx, (grp, fn, desc) in enumerate(func_data, start=1):
        cells = func_table.rows[idx].cells
        cells[0].text = grp
        cells[1].text = fn
        cells[2].text = desc
        bg = "F7F9FC" if idx % 2 == 1 else "FFFFFF"
        for c_idx, c in enumerate(cells):
            set_cell_background(c, bg)
            set_cell_margins(c, top=90, bottom=90, left=90, right=90)
            p = c.paragraphs[0]
            if c_idx == 0:
                p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # d. CÁC RÀNG BUỘC TRONG XỬ LÝ
    # -------------------------------------------------------------
    add_sec_heading("d. Các ràng buộc trong xử lý (Business Rules & Constraints)")

    add_sub_heading("1. Ràng buộc Đăng ký, Đăng nhập & Phiên làm việc (Session)")
    p = doc.add_paragraph()
    add_bullet(p, "• Đăng ký tài khoản:", "Username dài từ 3 đến 50 ký tự, tự động chuyển về chữ thường, kiểm tra trùng lặp trong cơ sở dữ liệu; Họ tên bắt buộc, tối đa 100 ký tự; Mật khẩu tối thiểu 6 ký tự và phải khớp tuyệt đối với Xác nhận mật khẩu. Mật khẩu được mã hóa SHA-256 trước khi lưu.")
    p = doc.add_paragraph()
    add_bullet(p, "• Đăng nhập & Trạng thái khóa:", "Username và mật khẩu phải chính xác. Tài khoản có IsLocked = true sẽ bị chặn hoàn toàn và thông báo tài khoản bị khóa do vi phạm. Tài khoản có ApprovedReportCount == 2 sẽ hiển thị cảnh báo vi phạm khi đăng nhập.")
    p = doc.add_paragraph()
    add_bullet(p, "• Cấu hình Session:", "Phiên làm việc hết hạn sau 2 giờ không hoạt động (IdleTimeout = 2 hours). Cookie phiên cấu hình HttpOnly = true, IsEssential = true để bảo vệ an toàn trước tấn công XSS.")

    add_sub_heading("2. Ràng buộc Hồ sơ cá nhân & Đổi mật khẩu")
    p = doc.add_paragraph()
    add_bullet(p, "• Hồ sơ cá nhân:", "Họ tên không được để trống hoặc chỉ có khoảng trắng. File ảnh đại diện upload chỉ chấp nhận các đuôi mở rộng: .jpg, .jpeg, .png, .gif, .svg; file được đổi tên bằng GUID ngẫu nhiên để chống ghi đè trên máy chủ.")
    p = doc.add_paragraph()
    add_bullet(p, "• Đổi mật khẩu:", "Người dùng bắt buộc phải nhập đúng mật khẩu hiện tại (so sánh chuỗi băm SHA-256); mật khẩu mới phải từ 6 ký tự trở lên và phải trùng khớp với xác nhận mật khẩu mới.")

    add_sub_heading("3. Ràng buộc Đăng bài viết, Bình luận & Bộ lọc từ cấm")
    p = doc.add_paragraph()
    add_bullet(p, "• Đăng bài viết:", "Bài viết bắt buộc phải có nội dung chữ (tối đa 1000 ký tự) HOẶC có hình ảnh đính kèm (hoặc cả hai). Hệ thống từ chối đăng bài hoàn toàn trống.")
    p = doc.add_paragraph()
    add_bullet(p, "• Bộ lọc từ cấm (Banned Words):", "Nội dung bài viết và bình luận tự động được kiểm duyệt với danh sách BannedWords bằng biểu thức chính quy Regex theo biên từ (?i)(?:\\s|^)word(?:\\s|[.,!?;:]|$). Nếu phát hiện, hệ thống từ chối lưu và trả về danh sách từ vi phạm để người dùng chỉnh sửa.")
    p = doc.add_paragraph()
    add_bullet(p, "• Bình luận & Phản hồi:", "Nội dung bình luận không được để trống (tối đa 500 ký tự); bài viết cha phải tồn tại; trường ParentCommentId phải trỏ chính xác vào ID bình luận cha nếu là bình luận phản hồi.")

    add_sub_heading("4. Ràng buộc Quyền xóa nội dung (Authorization)")
    p = doc.add_paragraph()
    add_bullet(p, "• Xóa bài viết:", "Chỉ tác giả bài viết (post.UserId == currentUser.UserId) hoặc Quản trị viên (currentUser.IsAdmin == true) mới có quyền xóa bài viết.")
    p = doc.add_paragraph()
    add_bullet(p, "• Xóa bình luận:", "Chỉ tác giả bình luận, chủ sở hữu bài viết hoặc Admin mới có quyền xóa bình luận. Các tài khoản khác bị từ chối.")

    add_sub_heading("5. Ràng buộc Bày tỏ cảm xúc (Reaction)")
    p = doc.add_paragraph()
    add_bullet(p, "• Tính duy nhất:", "Hỗ trợ 6 loại cảm xúc: Like, Love, Haha, Wow, Sad, Angry. Mỗi người dùng chỉ có duy nhất 1 bản ghi cảm xúc trên mỗi bài viết hoặc trên mỗi bình luận.")
    p = doc.add_paragraph()
    add_bullet(p, "• Cơ chế Toggle:", "Chọn lại cùng loại cảm xúc đã thả -> Gỡ bỏ cảm xúc. Chọn loại cảm xúc khác -> Cập nhật sang loại mới. Tự động tính tổng và trích xuất Top 3 emoji phổ biến nhất.")

    add_sub_heading("6. Ràng buộc Quản lý bạn bè & Nhắn tin trò chuyện (Chat)")
    p = doc.add_paragraph()
    add_bullet(p, "• Kết bạn:", "Không được kết bạn với chính mình; không kết bạn với tài khoản Admin; không tạo thêm lời mời nếu đã tồn tại mối quan hệ giữa 2 người; chỉ người nhận (ReceiverId) mới có quyền chấp nhận lời mời.")
    p = doc.add_paragraph()
    add_bullet(p, "• Nhắn tin (Chat):", "Hai người bắt buộc phải có quan hệ bạn bè đã được chấp nhận (IsAccepted == true) mới được phép nhắn tin và xem lịch sử chat của nhau. Nội dung tin nhắn không được rỗng (phải có chữ, ảnh hoặc bài viết chia sẻ).")
    p = doc.add_paragraph()
    add_bullet(p, "• Chia sẻ bài viết:", "Khóa ngoại SharedPostId trong bảng Messages được cấu hình OnDelete(DeleteBehavior.SetNull) để khi bài viết gốc bị xóa thì tin nhắn chia sẻ vẫn được bảo toàn.")

    add_sub_heading("7. Ràng buộc Báo cáo vi phạm & Cơ chế xử phạt bậc thang")
    p = doc.add_paragraph()
    add_bullet(p, "• Báo cáo bài viết:", "Bài viết phải tồn tại; lý do báo cáo bắt buộc (tối đa 500 ký tự); mỗi người dùng chỉ được báo cáo 1 bài viết tối đa 1 lần để chống spam.")
    p = doc.add_paragraph()
    add_bullet(p, "• Xử phạt khi Admin duyệt báo cáo (ApproveReport):", "Xóa bài viết vi phạm và tăng ApprovedReportCount của tác giả thêm 1. Quy chế xử lý bậc thang:")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.add_run("- Vi phạm lần 1: Ghi nhận vi phạm vào hệ thống.\n"
              "- Vi phạm lần 2: Cho phép đăng nhập nhưng hiển thị cảnh báo đỏ trên trang chủ.\n"
              "- Vi phạm lần 3 trở lên: Hệ thống tự động đặt cờ IsLocked = true, khóa tài khoản vĩnh viễn.")

    add_sub_heading("8. Ràng buộc Quản trị viên (Admin Rules)")
    p = doc.add_paragraph()
    add_bullet(p, "• Bảo vệ tài khoản Admin:", "Admin không được phép tự khóa hoặc tự xóa tài khoản của chính mình.")
    p = doc.add_paragraph()
    add_bullet(p, "• Xóa người dùng an toàn:", "Khi xóa tài khoản thành viên (DeleteUser), hệ thống chủ động xóa cascade tuần tự: Friendships, PostLikes, CommentLikes, Comments, PostReports, Posts trước khi xóa User để tránh vi phạm khóa ngoại DeleteBehavior.Restrict của SQL Server.")
    p = doc.add_paragraph()
    add_bullet(p, "• Quản lý từ cấm:", "Từ cấm không được để trống, không được trùng lặp với từ đã có trong hệ thống (so sánh không phân biệt chữ hoa/thường).")

    # -------------------------------------------------------------
    # e. PROCESS XỬ LÝ (RIÊNG / CHUNG)
    # -------------------------------------------------------------
    add_sec_heading("e. Process xử lý (Quy trình chung và quy trình riêng)")

    add_sub_heading("1. Quy trình xử lý chung của hệ thống")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng truy cập website qua trình duyệt.\n"
              "2. Hệ thống kiểm tra trạng thái Session đăng nhập:\n"
              "   - Nếu chưa đăng nhập: Cho phép xem các trang công khai (About, Privacy) hoặc yêu cầu Đăng ký / Đăng nhập đối với các tính năng cá nhân.\n"
              "   - Nếu đã đăng nhập: Phân luồng vai trò (Admin chuyển vào /Admin, User thông thường chuyển vào /Home/Index).\n"
              "3. Người dùng nhập thông tin và gửi yêu cầu đăng nhập/đăng ký.\n"
              "4. Server xác thực thông tin tài khoản và kiểm tra trạng thái khóa (IsLocked). Nếu tài khoản có 2 lần vi phạm đã duyệt thì gán cảnh báo vi phạm.\n"
              "5. Người dùng thao tác chức năng (đăng bài, bình luận, like, kết bạn, chat, báo cáo, quản trị...).\n"
              "6. Server kiểm tra tính hợp lệ của Session, quyền hạn (Authorization), ràng buộc dữ liệu nghiệp vụ và lọc từ cấm.\n"
              "7. Nếu không hợp lệ: Trả về thông báo lỗi chi tiết (JSON success = false, message).\n"
              "8. Nếu hợp lệ: Thực thi lưu dữ liệu vào SQL Server qua EF Core, sau đó phát sóng sự kiện qua SignalR đến các client trực tuyến.\n"
              "9. Giao diện người dùng cập nhật tức thời theo thời gian thực (Real-time) mà không cần tải lại trang.\n"
              "10. Người dùng tiếp tục tương tác hoặc kết thúc phiên bằng chức năng Đăng xuất.")

    add_sub_heading("2. Quy trình riêng: Đăng bài viết mới")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng nhập nội dung văn bản và/hoặc chọn tệp hình ảnh đính kèm.\n"
              "2. Nhấn nút 'Đăng bài'. Trình duyệt gửi AJAX POST /Post/Create (FormData).\n"
              "3. Server xác thực Session đăng nhập và kiểm tra dữ liệu không được rỗng cả chữ và ảnh.\n"
              "4. Nếu có nội dung văn bản: Hệ thống truy vấn bảng BannedWords, áp dụng Regex quét từ cấm.\n"
              "   - Nếu phát hiện từ cấm: Trả về lỗi kèm danh sách từ ngữ vi phạm, yêu cầu sửa lại.\n"
              "5. Nếu có hình ảnh: Kiểm tra, tạo thư mục /wwwroot/uploads/posts/, sinh tên file GUID ngẫu nhiên và lưu file lên máy chủ.\n"
              "6. Khởi tạo đối tượng Post mới, gán thông tin tác giả, thời gian tạo và lưu vào Database.\n"
              "7. Server gọi SignalR Hub phát sự kiện 'ReceiveNewPost' tới toàn bộ client đang online.\n"
              "8. Trả về kết quả thành công, giao diện client tự động chèn bài viết mới lên đầu bảng tin.")

    add_sub_heading("3. Quy trình riêng: Bình luận và trả lời bình luận")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng nhập nội dung vào ô bình luận dưới bài viết hoặc chọn 'Trả lời' dưới bình luận khác (parentCommentId).\n"
              "2. Nhấn nút gửi bình luận (AJAX POST /Post/CreateComment).\n"
              "3. Server kiểm tra Session, kiểm tra bài viết gốc tồn tại và nội dung bình luận không được rỗng (tối đa 500 ký tự).\n"
              "4. Kiểm tra từ cấm qua danh sách BannedWords; nếu vi phạm thì từ chối lưu.\n"
              "5. Tạo đối tượng Comment mới (gán PostId, UserId, ParentCommentId nếu có) và lưu vào Database.\n"
              "6. Phát sự kiện SignalR 'ReceiveNewComment' truyền dữ liệu bình luận mới tới các client đang theo dõi bài viết.\n"
              "7. Giao diện client tự động hiển thị bình luận mới vào đúng khối bình luận gốc hoặc phân cấp dưới bình luận cha.")

    add_sub_heading("4. Quy trình riêng: Bày tỏ cảm xúc (Reactions)")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng chọn 1 trong 6 cảm xúc (Like, Love, Haha, Wow, Sad, Angry) trên bài viết hoặc bình luận.\n"
              "2. Client gửi AJAX POST /Post/ToggleLike hoặc /Post/ToggleCommentLike.\n"
              "3. Server tìm kiếm bản ghi cảm xúc hiện tại của người dùng đối với đối tượng:\n"
              "   - Chưa có cảm xúc: Tạo mới bản ghi với ReactionType được chọn.\n"
              "   - Chọn lại cùng loại cảm xúc đã thả: Xóa bản ghi (gỡ bỏ cảm xúc).\n"
              "   - Chọn loại cảm xúc khác: Cập nhật ReactionType mới cho bản ghi hiện có.\n"
              "4. Lưu thay đổi vào DB; gom nhóm các lượt tương tác để tính tổng số và trích xuất Top 3 biểu tượng phổ biến nhất.\n"
              "5. Phát sóng sự kiện SignalR 'UpdatePostReactions' hoặc 'UpdateCommentLikes'.\n"
              "6. Giao diện tất cả client đồng loạt cập nhật trạng thái nút like, màu sắc và số lượt hiển thị.")

    add_sub_heading("5. Quy trình riêng: Quản lý bạn bè")
    p = doc.add_paragraph()
    p.add_run("1. Gợi ý kết bạn: ViewComponent tự động đề xuất 5 người dùng ngẫu nhiên chưa có quan hệ kết bạn và không phải Admin.\n"
              "2. Gửi lời mời: Người dùng nhấn 'Kết bạn' -> Gửi POST /Friend/SendRequest?receiverId=... Server kiểm tra không gửi cho chính mình, không gửi cho Admin, chưa có quan hệ trước đó -> Tạo bản ghi Friendship với IsAccepted = false.\n"
              "3. Xử lý lời mời: Người nhận vào mục Bạn bè, xem danh sách lời mời đang chờ:\n"
              "   - Chấp nhận (AcceptRequest): Server cập nhật IsAccepted = true. Hai người chính thức trở thành bạn bè và mở khóa tính năng Chat.\n"
              "   - Từ chối / Hủy kết bạn (RemoveFriend): Server xóa bản ghi khỏi bảng Friendships.")

    add_sub_heading("6. Quy trình riêng: Nhắn tin và chia sẻ bài viết (Chat)")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng mở hộp thoại chat với một người trong danh sách bạn bè (hoặc nhấn 'Chia sẻ bài viết' từ bảng tin).\n"
              "2. Server kiểm tra điều kiện tiên quyết: Hai người bắt buộc phải là bạn bè đã chấp nhận (IsAccepted == true).\n"
              "3. Tải lịch sử trò chuyện qua GET /Chat/GetMessages?friendId=...\n"
              "4. Người dùng nhập tin nhắn chữ, chọn file ảnh hoặc gửi bài viết chia sẻ -> Gửi POST /Chat/SendMessage.\n"
              "5. Server kiểm tra tin nhắn không được rỗng cả chữ, ảnh và sharedPostId. Nếu có ảnh, lưu vào /uploads/chat/.\n"
              "6. Lưu đối tượng Message vào DB và phát sự kiện SignalR 'ReceiveDirectMessage' tới người nhận và người gửi.\n"
              "7. Giao diện hộp chat lập tức vẽ bong bóng tin nhắn (kèm ảnh hoặc thẻ bài viết chia sẻ) và cuộn xuống tin nhắn mới nhất.")

    add_sub_heading("7. Quy trình riêng: Báo cáo bài viết và cơ chế xử lý vi phạm bậc thang")
    p = doc.add_paragraph()
    p.add_run("1. Người dùng nhấn nút Báo cáo trên bài viết, nhập lý do vi phạm và gửi request.\n"
              "2. Server kiểm tra bài viết tồn tại, lý do không rỗng và kiểm tra người này chưa từng báo cáo bài viết đó trước đây -> Lưu bản ghi PostReport.\n"
              "3. Quản trị viên (Admin) truy cập Dashboard xem danh sách các báo cáo chờ duyệt.\n"
              "4. Admin đánh giá báo cáo:\n"
              "   - Bỏ qua (DismissReport): Xóa bản ghi PostReport khỏi DB, không xử phạt tác giả bài viết.\n"
              "   - Duyệt vi phạm (ApproveReport): Xóa bài viết vi phạm, đồng thời tăng ApprovedReportCount của tác giả thêm 1.\n"
              "5. Hệ thống kiểm tra số lần vi phạm tích lũy của tác giả:\n"
              "   - Lần 1: Ghi nhận vi phạm vào hồ sơ tài khoản.\n"
              "   - Lần 2: Gán cảnh báo vi phạm; khi người dùng đăng nhập sẽ xuất hiện thông báo cảnh báo đỏ trên trang chủ.\n"
              "   - Từ lần 3 trở lên: Hệ thống tự động đặt cờ IsLocked = true, khóa tài khoản vĩnh viễn và chặn mọi phiên đăng nhập tiếp theo.")

    add_sub_heading("8. Quy trình riêng: Quản trị người dùng của Admin")
    p = doc.add_paragraph()
    p.add_run("1. Admin mở trang /Admin/Users để xem toàn bộ danh sách thành viên.\n"
              "2. Đổi mật khẩu (ChangeUserPassword): Admin nhập mật khẩu mới cho user (>=6 ký tự) -> Server băm SHA-256 và cập nhật PasswordHash.\n"
              "3. Khóa / Mở khóa (ToggleLockUser): Admin bấm nút tương ứng -> Server kiểm tra không khóa chính mình -> Đảo ngược trạng thái user.IsLocked.\n"
              "4. Xóa tài khoản vĩnh viễn (DeleteUser): Server kiểm tra không xóa chính mình -> Thực thi xóa cascade thủ công các liên kết: Friendships, PostLikes, CommentLikes, Comments, PostReports, Posts -> Sau cùng xóa bản ghi User.")

    add_sub_heading("9. Quy trình riêng: Quản lý từ cấm (Banned Words)")
    p = doc.add_paragraph()
    p.add_run("1. Admin truy cập /Admin/BannedWords để xem danh sách từ cấm đang áp dụng.\n"
              "2. Thêm từ cấm (AddBannedWord): Admin nhập từ mới -> Server kiểm tra không rỗng và không trùng lặp (không phân biệt hoa/thường) -> Lưu bản ghi BannedWord mới. Từ cấm lập tức có hiệu lực kiểm duyệt cho các bài viết/bình luận tiếp theo.\n"
              "3. Xóa từ cấm (DeleteBannedWord): Admin nhấn xóa bên cạnh từ cấm -> Server xóa bản ghi khỏi DB -> Bỏ chặn từ ngữ này trên hệ thống.")

    add_sub_heading("10. Quy trình riêng: Khởi tạo dữ liệu mẫu khi khởi động hệ thống (Database Seeding)")
    p = doc.add_paragraph()
    p.add_run("1. Ứng dụng khởi chạy trong Program.cs, tạo IServiceScope để lấy SocialDbContext.\n"
              "2. Gọi phương thức SeedData.SeedDatabase(context).\n"
              "3. Tự động áp dụng Migration (context.Database.Migrate()) và kiểm tra cấu trúc bảng Messages trong SQL Server.\n"
              "4. Kiểm tra và chèn 13 từ cấm mặc định nếu bảng BannedWords đang rỗng.\n"
              "5. Kiểm tra và tạo tài khoản Admin mặc định (username: admin / mật khẩu: admin123) nếu chưa có.\n"
              "6. Khởi tạo 7 người dùng mẫu (vietanh, lanhuong, quanghuy, tantai, hanh, minhtuan, phuongthao với mật khẩu mặc định 123456).\n"
              "7. Tạo sẵn các bài viết mẫu (có ảnh Unsplash và text), bình luận phân cấp, lượt thích đa dạng loại cảm xúc và quan hệ bạn bè mẫu để hệ thống sẵn sàng sử dụng ngay.")

    # Save document
    doc.save(docx_path)
    print(f"Successfully generated docx at: {docx_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\PC\Desktop\SocialWebNetWorking"
    docx_file = os.path.join(out_dir, "TD01_NhomXX.docx")
    build_document(docx_file)
